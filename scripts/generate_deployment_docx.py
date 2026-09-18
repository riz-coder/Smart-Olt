from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "OptiVerse_Nexecode_Public_VPS_Deployment_Guide.md"
OUTPUT = ROOT / "docs" / "OptiVerse_Nexecode_Public_VPS_Deployment_Guide.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=160, bottom=120, end=160):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field = OxmlElement("w:instrText")
    field.set(qn("xml:space"), "preserve")
    field.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, field, separate, end))


def add_inline(paragraph, text, *, bold=False):
    pieces = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(18, 75, 115)
        elif piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        else:
            paragraph.add_run(piece).bold = bold


def add_code_box(document, code, language=""):
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF3F8")
    set_cell_margins(cell)
    if language:
        label = cell.paragraphs[0]
        label.paragraph_format.space_after = Pt(3)
        run = label.add_run("COMMAND" if language in {"bash", "sh", "powershell"} else "CONFIGURATION")
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(32, 99, 155)
    else:
        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    paragraph = cell.add_paragraph() if language else cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(code.rstrip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(26, 35, 48)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for name, size, color in (
        ("Title", 24, RGBColor(18, 64, 98)),
        ("Heading 1", 16, RGBColor(18, 83, 128)),
        ("Heading 2", 13, RGBColor(35, 99, 150)),
        ("Heading 3", 11, RGBColor(43, 110, 160)),
    ):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = color

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("OptiVerse Production Guide  |  Page ")
    add_field(footer, "PAGE")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_code = False
    language = ""
    code_lines = []
    first_heading = True
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                add_code_box(document, "\n".join(code_lines), language)
                in_code = False
                language = ""
                code_lines = []
            else:
                in_code = True
                language = line[3:].strip().lower()
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if first_heading and level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run(text)
                subtitle = document.add_paragraph()
                subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = subtitle.add_run("Production deployment, Linux verification and tenant provisioning")
                run.italic = True
                run.font.color.rgb = RGBColor(90, 105, 120)
                document.add_paragraph()
                first_heading = False
            else:
                document.add_heading(text, level=level)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|(?:\s*:?-+:?\s*\|)+$", lines[index + 1]):
            headers = [item.strip() for item in line.strip("|").split("|")]
            rows = []
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([item.strip() for item in lines[index].strip("|").split("|")])
                index += 1
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for column, value in enumerate(headers):
                set_cell_shading(table.rows[0].cells[column], "D9EAF7")
                add_inline(table.rows[0].cells[column].paragraphs[0], value, bold=True)
            for values in rows:
                cells = table.add_row().cells
                for column, value in enumerate(values[:len(headers)]):
                    add_inline(cells[column].paragraphs[0], value)
            document.add_paragraph()
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(1))
        elif numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(1))
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
        index += 1

    document.core_properties.title = "OptiVerse Public Deployment, Verification and Control Panel Guide"
    document.core_properties.subject = "Production deployment and operations handbook"
    document.core_properties.author = "OptiVerse"
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
